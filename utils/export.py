from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from database.queries import leads as leads_q, sources
from utils.fmt import to_local

HEADERS = ["ID", "Sana", "Ism", "Telefon", "Telegram", "Kurs", "Turi", "Manba", "Status",
           "Uchrashuv", "Rad sababi", "Izohlar", "Yangilangan"]


async def leads_to_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Arizalar"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        tg = f"@{r['username']}" if r.get("username") else (str(r["telegram_id"]) if r.get("telegram_id") else "")
        ws.append([
            r["id"], to_local(r["created_at"]), r["name"], r["phone"], tg, r.get("course_title") or "",
            leads_q.TYPE_LABELS.get(r["type"], r["type"]), await sources.title_of(r.get("source")),
            leads_q.status_label(r["status"]), r.get("meeting_at") or "", r.get("reject_reason") or "",
            r.get("notes") or "", to_local(r["updated_at"]),
        ])
    widths = [6, 17, 22, 16, 18, 24, 18, 18, 22, 18, 24, 40, 17]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
